from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.game import Game


def _formatar_data(timestamp: int | None) -> str:
    """Converte timestamp Unix da IGDB para data legível (DD/MM/AAAA)."""
    if timestamp is None:
        return ""
    data = datetime.fromtimestamp(timestamp, tz=timezone.utc)
    return data.strftime("%d/%m/%Y")


def export_games_to_excel(db: Session, output_path: str) -> None:
    """Exporta todos os jogos do banco para um arquivo Excel formatado."""
    jogos = db.query(Game).order_by(Game.name).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Jogos"

    colunas = ["Nome", "Data de Lançamento", "Nota", "Gêneros", "Descrição"]
    ws.append(colunas)

    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    for coluna_idx in range(1, len(colunas) + 1):
        celula = ws.cell(row=1, column=coluna_idx)
        celula.font = cabecalho_fonte
        celula.fill = cabecalho_fundo
        celula.alignment = Alignment(horizontal="center", vertical="center")

    for jogo in jogos:
        ws.append([
            jogo.name,
            _formatar_data(jogo.release_date),
            round(jogo.rating, 1) if jogo.rating else "",
            jogo.genres or "",
            jogo.summary or "",
        ])

    larguras = {"A": 35, "B": 18, "C": 10, "D": 25, "E": 60}
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura

    for row in ws.iter_rows(min_row=2, max_col=len(colunas)):
        for celula in row:
            celula.alignment = Alignment(vertical="top", wrap_text=(celula.column_letter == "E"))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(jogos) + 1}"

    wb.save(output_path)